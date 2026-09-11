#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""solve_b.py — CUMCM 2026 B 题「无线电干扰源的快速自动定位与清除」求解器。

四个子问题
----------
Q1  交会定位法 → 多边形定位区域直径算法 + 「以直径为直径的圆能否覆盖该区域」判定。
Q2  第二检测点选择策略 + 候选区域（最优判据：交会角 ∠≈90°，即 G_est 处 Thales 切线）。
Q3  全向干扰源：机器狗自动搜索 / 定位 / 清除策略 + 本地模拟器演练统计。
Q4  混合全向 + 定向：策略与算法扩展 + 演练统计。

设计要点（对齐 Modeling-Harness 铁律）
-------------------------------------
* 坐标系：目标区域 R=1800 m 圆域，原点在圆心，x 正东、y 正北。
* **定位与逼近只使用测向机可观测的量**（示向度 + 信号相对强度）；
  真实源坐标仅由模拟器内部用于判定「是否进入光学/激光作用距离」，
  策略代码中不出现任何真值导向（无 ground-truth 泄漏、无占位兜底）。
* 随机种子固定 42；演练每问 30 个随机案例（≥5，满足多种子要求）。
* 时间口径：单次检测 5 s（附录2(5)「切换频道、调天线、稳定读数共 5 秒」，
  频道切换费已计入，不重复计 1 s）；光学定位 3 s；激光清除 2 s；机器狗 5 m/s。
"""
from __future__ import annotations

import json
import math
import os
import random

import numpy as np

# ----------------------------------------------------------------------
# 全局常量
# ----------------------------------------------------------------------
SEED = 42                                   # 随机种子（铁律：固定 42）
R_AREA = 1800.0                             # 目标区域半径 / m
EPS_BEARING = 1.0                           # 示向度误差界 / °（全局 [−1,1]）
R_REC_MIN, R_REC_MAX = 1000.0, 1500.0       # 有效接收半径范围 / m
R_REC_NOM = 0.5 * (R_REC_MIN + R_REC_MAX)   # 名义接收半径（用于由强度估距）/ m
N_SRC_MIN, N_SRC_MAX = 10, 16               # 干扰源个数范围
N_CHANNEL = 20                              # 频道总数 1..20

T_DETECT = 5.0                              # 单次检测（含切换频道）/ s
T_OPTIC = 3.0                               # 光学精确定位 / s
T_LASER = 2.0                               # 激光清除 / s
D_OPTIC = 20.0                              # 光学定位距离阈值 / m
D_DIRECT = 5.0                              # 信号过强直读距离阈值 / m
V_DOG = 5.0                                 # 机器狗速度 / m·s⁻¹

# 覆盖扫描参数（同心环）
COVER_RADII_OMNI = (450.0, 1350.0)          # 全向源覆盖环：任意点距最近环 ≤450 m
COVER_RADII_DIR = (450.0, 1350.0, 2000.0)   # 另加「源分布环外侧」环以覆盖朝外定向源
COVER_STEP = 500.0                          # 内环（r≤1000 m）检测点弧长间隔 / m
COVER_STEP_MID = 900.0                      # 次外环（1000<r≤1500 m）检测点弧长间隔 / m
COVER_STEP_DIR = 1050.0                     # 外侧环（r>1500 m）检测点弧长间隔 / m


# ----------------------------------------------------------------------
# 基础几何工具
# ----------------------------------------------------------------------
def _dir(angle_deg: float):
    a = math.radians(angle_deg)
    return math.cos(a), math.sin(a)


def _bearing(x1: float, y1: float, x2: float, y2: float) -> float:
    """从 (x1,y1) 指向 (x2,y2) 的方位角 / ° ∈ [0,360)。"""
    return math.degrees(math.atan2(y2 - y1, x2 - x1)) % 360.0


def _bearing_diff(b1: float, b2: float) -> float:
    """两个方位角之夹角 / ° ∈ [0,180]（即示向度之差的绝对值）。"""
    return abs(((b1 - b2 + 180.0) % 360.0) - 180.0)


def _clamp(v, lo, hi):
    return max(lo, min(hi, v))


def _polygon_area(verts) -> float:
    if len(verts) < 3:
        return 0.0
    s = 0.0
    n = len(verts)
    for i in range(n):
        x1, y1 = verts[i]
        x2, y2 = verts[(i + 1) % n]
        s += x1 * y2 - x2 * y1
    return abs(s) / 2.0


# ----------------------------------------------------------------------
# Q1：交会定位多边形与直径
# ----------------------------------------------------------------------
def _wedge_halfplanes(sx: float, sy: float, theta_deg: float, eps: float):
    """楔形 W={P: arg(P−S) ∈ [θ−ε, θ+ε]} 的两条半平面 (a,b,c,sense)。

    sense='ge' 表示 a·x+b·y ≥ c；'le' 表示 ≤ c。由叉积方向（逆时针为正）判定。
    """
    out = []
    for ang in (theta_deg - eps, theta_deg + eps):
        dx, dy = _dir(ang)
        a, b = -dy, dx                      # cross(d, P−S) = a·Px + b·Py − (a·Sx + b·Sy)
        c = a * sx + b * sy
        out.append((a, b, c))
    return [(out[0][0], out[0][1], out[0][2], "ge"),
            (out[1][0], out[1][1], out[1][2], "le")]


def _clip_halfplane(poly, a, b, c, sense):
    """Sutherland–Hodgman：用半平面裁剪凸多边形。"""
    if not poly:
        return []

    def inside(p):
        val = a * p[0] + b * p[1]
        return val >= c - 1e-9 if sense == "ge" else val <= c + 1e-9

    def inter(p, q):
        fp = a * p[0] + b * p[1] - c
        fq = a * q[0] + b * q[1] - c
        t = fp / (fp - fq) if (fp - fq) != 0 else 0.0
        return (p[0] + t * (q[0] - p[0]), p[1] + t * (q[1] - p[1]))

    out = []
    n = len(poly)
    for i in range(n):
        p, q = poly[i], poly[(i + 1) % n]
        pin, qin = inside(p), inside(q)
        if pin:
            out.append(p)
        if pin != qin:
            out.append(inter(p, q))
    return out


def location_region(points, bearings, eps=EPS_BEARING, bbox=8000.0):
    """定位区域 = 各楔形（示向度 ± ε）之交集（凸多边形）。返回顶点列表。"""
    poly = [(-bbox, -bbox), (bbox, -bbox), (bbox, bbox), (-bbox, bbox)]
    for (sx, sy), th in zip(points, bearings):
        for (a, b, c, sense) in _wedge_halfplanes(sx, sy, th, eps):
            poly = _clip_halfplane(poly, a, b, c, sense)
            if len(poly) < 3:
                return poly
    # 去重
    seen, out = set(), []
    for p in poly:
        key = (round(p[0], 6), round(p[1], 6))
        if key not in seen:
            seen.add(key)
            out.append(p)
    return out


def _convex_hull(pts):
    pts = sorted(set((round(x, 9), round(y, 9)) for x, y in pts))
    if len(pts) <= 2:
        return pts

    def cross(o, a, b):
        return (a[0] - o[0]) * (b[1] - o[1]) - (a[1] - o[1]) * (b[0] - o[0])

    lower = []
    for p in pts:
        while len(lower) >= 2 and cross(lower[-2], lower[-1], p) <= 0:
            lower.pop()
        lower.append(p)
    upper = []
    for p in reversed(pts):
        while len(upper) >= 2 and cross(upper[-2], upper[-1], p) <= 0:
            upper.pop()
        upper.append(p)
    return lower[:-1] + upper[:-1]


def polygon_diameter(verts):
    """凸多边形直径（任意两点最大距离）= 凸包上 O(n²) 暴力（等价旋转卡壳）。"""
    hull = _convex_hull(verts)
    best, pair = 0.0, None
    for i in range(len(hull)):
        for j in range(i + 1, len(hull)):
            d = math.dist(hull[i], hull[j])
            if d > best:
                best, pair = d, (hull[i], hull[j])
    return best, pair, hull


def min_enclosing_circle(verts):
    """最小包围圆（Welzl 随机增量，期望 O(n)）。返回 (center, radius)。"""
    pts = [tuple(v) for v in verts]
    if not pts:
        return (0.0, 0.0), 0.0
    if len(pts) == 1:
        return pts[0], 0.0
    rng = random.Random(SEED)

    def circle_from2(a, b):
        c0 = ((a[0] + b[0]) / 2, (a[1] + b[1]) / 2)
        return c0, math.dist(a, b) / 2

    def circle_from3(a, b, c):
        ax, ay = a
        bx, by = b
        cx, cy = c
        d = 2 * (ax * (by - cy) + bx * (cy - ay) + cx * (ay - by))
        if abs(d) < 1e-12:
            pair = max([(a, b), (b, c), (a, c)], key=lambda t: math.dist(t[0], t[1]))
            return circle_from2(*pair)
        ux = ((ax * ax + ay * ay) * (by - cy) + (bx * bx + by * by) * (cy - ay)
              + (cx * cx + cy * cy) * (ay - by)) / d
        uy = ((ax * ax + ay * ay) * (cx - bx) + (bx * bx + by * by) * (ax - cx)
              + (cx * cx + cy * cy) * (bx - ax)) / d
        c0 = (ux, uy)
        return c0, math.dist(c0, a)

    def in_circle(p, c, r):
        return math.dist(p, c) <= r + 1e-9

    P = pts[:]
    rng.shuffle(P)
    c, r = P[0], 0.0
    for i in range(1, len(P)):
        if not in_circle(P[i], c, r):
            c, r = P[i], 0.0
            for j in range(i):
                if not in_circle(P[j], c, r):
                    c, r = circle_from2(P[i], P[j])
                    for k in range(j):
                        if not in_circle(P[k], c, r):
                            c, r = circle_from3(P[i], P[j], P[k])
    return c, r


def diameter_circle_covers(verts):
    """以「定位区域直径」为直径的圆能否覆盖该区域？

    凸集 ⇒ 只需检验全部顶点是否落在直径圆内。返回
    (covers, diam, center_diam, radius_diam, mec_center, mec_radius)。
    """
    diam, pair, hull = polygon_diameter(verts)
    if pair is None:
        return False, 0.0, (0.0, 0.0), 0.0, (0.0, 0.0), 0.0
    center_d = ((pair[0][0] + pair[1][0]) / 2, (pair[0][1] + pair[1][1]) / 2)
    radius_d = diam / 2
    covers = all(math.dist(v, center_d) <= radius_d + 1e-6 for v in hull)
    mec_c, mec_r = min_enclosing_circle(hull)
    return covers, diam, center_d, radius_d, mec_c, mec_r


def problem1_demo():
    """Q1 演示：规范交会案例 + 直径圆覆盖判定 + 反例。"""
    print("=" * 68)
    print("问题1：交会定位多边形直径算法")
    cases = {
        "A 正交交会（φ≈90°）": ([(0.0, 0.0), (3000.0, 0.0)], (1500.0, 1500.0)),
        "B 钝角交会（φ≈142°）": ([(0.0, 0.0), (3000.0, 0.0)], (1200.0, 500.0)),
        "C 三站交会": ([(0.0, 0.0), (3000.0, 0.0), (1500.0, 2600.0)], (1500.0, 1200.0)),
    }
    results = {}
    for name, (stations, G) in cases.items():
        bears = [_bearing(sx, sy, G[0], G[1]) for (sx, sy) in stations]
        reg = location_region(stations, bears)
        if len(reg) < 3:
            print(f"  {name}: 定位区域退化（空/无界）")
            results[name] = {"degenerate": True}
            continue
        covers, diam, cd, rd, mec_c, mec_r = diameter_circle_covers(reg)
        area = _polygon_area(reg)
        phi = _bearing_diff(bears[0], bears[1])
        results[name] = {"n_vertices": len(reg), "crossing_angle_deg": phi,
                         "diameter_m": diam, "area_m2": area,
                         "diam_circle_r_m": rd, "mec_r_m": mec_r,
                         "covers": covers}
        print(f"  {name}: 交会角={phi:6.2f}°  顶点数={len(reg)}  "
              f"直径={diam:8.3f} m  面积={area:12.3f} m²")
        print(f"      直径圆半径={rd:8.3f} m  最小包围圆半径={mec_r:8.3f} m  "
              f"直径圆覆盖={covers}")

    # 反例：锐角三角形（第三顶点落在以最长边为直径的圆之外）
    tri = [(0.0, 0.0), (10.0, 0.0), (5.0, 6.0)]
    covers, diam, cd, rd, mec_c, mec_r = diameter_circle_covers(tri)
    print(f"  反例（锐角三角形）: 直径={diam:.4f} m  直径圆半径={rd:.4f} m  "
          f"最小包围圆半径={mec_r:.4f} m  直径圆覆盖={covers}")
    print("  结论：以定位区域直径为直径的圆『不一定』能覆盖该区域。")
    print("        由 Thales 定理：凸多边形被直径圆覆盖 ⟺ 其余各顶点对直径端点")
    print("        的张角均 ≥90°（多边形对最长弦不呈『锐角』）。反例中第三顶点对")
    print("        直径端点张角 79.6°<90°，落在直径圆外，故覆盖失败；")
    print("        保证覆盖应改用最小包围圆（Welzl，半径 mec_r）。")
    results["反例(锐角三角形)"] = {"diameter_m": diam, "diam_circle_r_m": rd,
                                  "mec_r_m": mec_r, "covers": covers}
    return results


# ----------------------------------------------------------------------
# Q2：第二检测点选择策略
# ----------------------------------------------------------------------
def crossing_angle(S1, b1, S2, b2):
    """两示向度在交点处的交会角 / °（= 两方位向之夹角，0..180）。"""
    return _bearing_diff(b1, b2)


def perpendicular_second_point(P1, b1, d_est, offset):
    """过估计源 G_est=P1+d_est·dir(b1) 且垂直于 P1→G_est 的直线上的点。

    该直线上任意点到 G_est 的方位与 b1 夹角恒为 90°（Thales 切线）。
    """
    gx = P1[0] + d_est * math.cos(math.radians(b1))
    gy = P1[1] + d_est * math.sin(math.radians(b1))
    px = gx + offset * math.cos(math.radians(b1 + 90.0))
    py = gy + offset * math.sin(math.radians(b1 + 90.0))
    return (px, py)


def candidate_second_point_region(P1, b1, d_est, phi_min=60.0, r_rec=R_REC_MAX,
                                  n_r=17, d_deg=5.0):
    """第二检测点候选区域（极坐标网格）。

    以估计源 G_est 为极点：候选点 P 满足
      (a) 交会角 φ = ∠(b1, P→G_est) ∈ [φ_min, 180−φ_min]（近正交最优）；
      (b) 0 < |P−G_est| ≤ r_rec（可接收到信号）。
    评分：主判据 |φ−90°| 最小；次判据定位区域直径最小（进一步减小误差区域）。
    返回 (grid, best) 其中 grid=[(x,y,φ)], best=(x,y,φ,diam) 或 (grid,None)。
    """
    G = (P1[0] + d_est * math.cos(math.radians(b1)),
         P1[1] + d_est * math.sin(math.radians(b1)))
    grid = []
    for rr in np.linspace(0.2 * r_rec, r_rec, n_r):
        for aa in np.arange(0.0, 360.0, d_deg):
            P = (G[0] + rr * math.cos(math.radians(aa)),
                 G[1] + rr * math.sin(math.radians(aa)))
            phi = _bearing_diff(b1, _bearing(P[0], P[1], G[0], G[1]))
            if phi_min <= phi <= 180.0 - phi_min:
                grid.append((P[0], P[1], phi))
    if not grid:
        return [], None
    scored = []
    for (x, y, phi) in grid:
        b2 = _bearing(x, y, G[0], G[1])
        reg = location_region([P1, (x, y)], [b1, b2])
        d = polygon_diameter(reg)[0] if len(reg) >= 3 else float("inf")
        scored.append((abs(phi - 90.0), d, x, y, phi))
    scored.sort(key=lambda t: (t[0], t[1]))
    _, d, x, y, phi = scored[0]
    return grid, (x, y, phi, d)


def problem2_demo():
    """Q2 演示：候选区域 + 最优判据 + 推荐点 + 偏移—精度权衡。"""
    print("=" * 68)
    print("问题2：第二个检测点选择策略")
    S1 = (-500.0, 0.0)
    theta1 = 20.0
    d_est = 1200.0
    G = (S1[0] + d_est * math.cos(math.radians(theta1)),
         S1[1] + d_est * math.sin(math.radians(theta1)))
    print(f"  已知：S1={S1}，示向度 θ1={theta1}°，源估计距离 d_est={d_est} m")
    print(f"  估计源位置 G_est=({G[0]:.1f},{G[1]:.1f})")
    grid, best = candidate_second_point_region(S1, theta1, d_est)
    print(f"  候选区域（φ∈[60°,120°]、|P−G_est|≤1500 m）网格点数={len(grid)}")
    if best:
        print(f"  网格最优：P2=({best[0]:.1f},{best[1]:.1f})  "
              f"交会角={best[2]:.2f}°  定位区域直径={best[3]:.2f} m")
    perp = perpendicular_second_point(S1, theta1, d_est, offset=d_est)
    phi_perp = _bearing_diff(theta1, _bearing(perp[0], perp[1], G[0], G[1]))
    print(f"  解析最优（过 G_est 的垂线，偏移=d_est）：P2=({perp[0]:.1f},{perp[1]:.1f})  "
          f"交会角={phi_perp:.2f}°")
    print("  偏移量—定位精度权衡（垂线上不同偏移）：")
    trade = []
    for off in (300.0, 600.0, 900.0, 1200.0, 1500.0):
        P2 = perpendicular_second_point(S1, theta1, d_est, offset=off)
        b2 = _bearing(P2[0], P2[1], G[0], G[1])
        reg = location_region([S1, P2], [theta1, b2])
        d = polygon_diameter(reg)[0] if len(reg) >= 3 else float("nan")
        a = _polygon_area(reg) if len(reg) >= 3 else float("nan")
        trade.append({"offset_m": off, "S2": [P2[0], P2[1]], "phi_deg": 90.0,
                      "region_diameter_m": d, "region_area_m2": a})
        print(f"    偏移={off:6.0f} m  S2=({P2[0]:7.1f},{P2[1]:7.1f})  "
              f"交会角={90.0:5.1f}°  定位区域直径={d:7.2f} m  面积={a:11.1f} m²")
    print("  结论：最优第二点位于过 G_est 且垂直于 S1→G_est 的直线上（交会角恒 90°），")
    print("        即 Thales 圆在 G_est 处的切线；沿该直线偏移越小定位区域越小，")
    print("        但偏移须使第二点仍在有效接收半径内且与第一点有足够基线")
    print("        （避免示向度误差相关），故推荐偏移 = d_est（对称平衡配置）。")
    return {"n_grid": len(grid),
            "grid_best": list(best[:3]) if best else None,
            "perp_recommend": [perp[0], perp[1]],
            "tradeoff": trade}


# ----------------------------------------------------------------------
# Q3/Q4：模拟器与机器狗策略
# ----------------------------------------------------------------------
class Source:
    __slots__ = ("sid", "x", "y", "channel", "r_rec", "kind", "heading")

    def __init__(self, sid, x, y, channel, r_rec, kind="omni", heading=None):
        self.sid = sid
        self.x = x
        self.y = y
        self.channel = channel
        self.r_rec = r_rec
        self.kind = kind                 # 'omni' 或 'dir'
        self.heading = heading           # 定向方向 / °；全向为 None

    def in_beam(self, px, py):
        if self.kind == "omni":
            return True
        ang = math.degrees(math.atan2(py - self.y, px - self.x)) % 360.0
        diff = abs(((ang - (self.heading or 0.0) + 180.0) % 360.0) - 180.0)
        return diff <= 90.0 + 1e-9


class Simulator:
    """本地演练模拟器（复现附录1/附录2 规则）。

    * 每源频道互异；信号互不影响；场强仅随距离衰减（用相对强度 d/r_rec 表示）。
    * 距离 ≤ r_rec 且在覆盖角内方可测向；距离 ≤5 m 且在内时信号过强无法测向。
    * 距离 ≤20 m 可光学精确定位（3 s）+ 激光清除（2 s）。
    """

    def __init__(self, seed=SEED, n_src=None, kind_mix=False, dir_frac=0.4):
        self.rng = random.Random(seed)
        n = n_src if n_src is not None else self.rng.randint(N_SRC_MIN, N_SRC_MAX)
        chans = list(range(1, N_CHANNEL + 1))
        self.rng.shuffle(chans)
        self.sources = []
        for i in range(n):
            r = R_AREA * math.sqrt(self.rng.random())
            a = self.rng.uniform(0.0, 2.0 * math.pi)
            kind, heading = "omni", None
            if kind_mix and self.rng.random() < dir_frac:
                kind, heading = "dir", self.rng.uniform(0.0, 360.0)
            self.sources.append(Source(i, r * math.cos(a), r * math.sin(a),
                                       chans[i],
                                       self.rng.uniform(R_REC_MIN, R_REC_MAX),
                                       kind, heading))
        self.rng.shuffle(self.sources)     # 打乱编号，避免与频道顺序相关
        self._err_cache = {}
        self.cleared = set()
        self.dog_pos = (0.0, 0.0)
        self.vtime = 0.0
        self.path_len = 0.0
        self.n_detect = 0
        self.n_move = 0

    # --- 执行原语 ---
    def move_to(self, x, y):
        d = math.dist(self.dog_pos, (x, y))
        self.path_len += d
        self.vtime += d / V_DOG
        self.dog_pos = (x, y)
        self.n_move += 1

    def is_channel_cleared(self, ch):
        for s in self.sources:
            if s.channel == ch and s.sid in self.cleared:
                return True
        return False

    def _bearing_error(self, src, px, py):
        key = (src.sid, round(px), round(py))
        if key not in self._err_cache:
            self._err_cache[key] = self.rng.uniform(-EPS_BEARING, EPS_BEARING)
        return self._err_cache[key]

    def detect(self, px, py, channel):
        """在 (px,py) 检测频道 channel，耗时 5 s。

        返回 (bearing_deg|None, rel|None)：rel=d/r_rec 为相对强度（[0,1]）。
        信号过强（d≤5）时返回 (None, 0.0)；无信号返回 (None, None)。
        """
        self.vtime += T_DETECT
        self.n_detect += 1
        best = None
        for src in self.sources:
            if src.channel != channel or src.sid in self.cleared:
                continue
            d = math.dist((px, py), (src.x, src.y))
            if d <= src.r_rec and src.in_beam(px, py):
                rel = d / src.r_rec
                if best is None or rel < best[1]:
                    best = (src, rel, d)
        if best is None:
            return None, None
        src, rel, d = best
        if d <= D_DIRECT:
            return None, 0.0
        ang = _bearing(px, py, src.x, src.y)
        return (ang + self._bearing_error(src, px, py)) % 360.0, rel

    def sweep_all_channels(self, px, py, channels=None):
        """在 (px,py) 依次扫描给定频道，返回 {channel: (bearing, rel)}。"""
        found = {}
        for ch in (channels if channels is not None else range(1, N_CHANNEL + 1)):
            b, rel = self.detect(px, py, ch)
            if b is not None and rel is not None and rel > 0.0:
                found[ch] = (b, rel)
        return found

    def try_clear(self, x, y):
        """在 (x,y) 尝试光学定位 + 清除最近的未清除源（距离 ≤20 m）。"""
        cands = [s for s in self.sources
                 if s.sid not in self.cleared
                 and math.dist((x, y), (s.x, s.y)) <= D_OPTIC]
        if not cands:
            return False
        src = min(cands, key=lambda s: math.dist((x, y), (s.x, s.y)))
        self.vtime += T_OPTIC + T_LASER
        self.cleared.add(src.sid)
        return True


# ---------- 交会与定位（仅用可观测量） ----------
def cross_fix(P1, b1, P2, b2):
    """两射线交会：射线 i 为 P_i + t·dir(b_i)，t≥0。返回交点或 None。"""
    d1, d2 = _dir(b1), _dir(b2)
    denom = d1[0] * d2[1] - d1[1] * d2[0]
    if abs(denom) < 1e-9:
        return None
    rx, ry = P2[0] - P1[0], P2[1] - P1[1]
    t = (rx * d2[1] - ry * d2[0]) / denom
    u = (rx * d1[1] - ry * d1[0]) / denom
    if t < -1e-6 or u < -1e-6:
        return None
    return (P1[0] + t * d1[0], P1[1] + t * d1[1])


def best_pair_fix(obs, phi_lo=25.0, phi_hi=155.0):
    """从观测集 {(P,b,rel)} 中取交会角最接近 90° 的两条做交会。"""
    best = None
    n = len(obs)
    for i in range(n):
        for j in range(i + 1, n):
            b1, b2 = obs[i][1], obs[j][1]
            phi = _bearing_diff(b1, b2)
            if not (phi_lo <= phi <= phi_hi):
                continue
            score = abs(phi - 90.0)
            if best is None or score < best[0]:
                est = cross_fix(obs[i][0], b1, obs[j][0], b2)
                if est is not None:
                    best = (score, est)
    return best[1] if best else None


def local_hunt(sim, est, ch, obs):
    """估计点附近环形搜索（含 ≤20 m 半径，便于直接光学清除）。

    已清除返回 True；重新捕获信号（追加 obs）返回 False。
    """
    for r in (15.0, 30.0, 45.0, 60.0, 90.0, 120.0):
        for k in range(8):
            a = math.radians(45.0 * k)
            p = (est[0] + r * math.cos(a), est[1] + r * math.sin(a))
            sim.move_to(*p)
            if sim.try_clear(*p):
                return True
            b, rel = sim.detect(p[0], p[1], ch)
            if b is not None and rel is not None and rel > 0.0:
                obs.append((p, b, rel))
                return False
    return False


def _homing(sim, ch, obs, max_step=500.0, n_iter=15):
    """归航逼近：沿当前示向度按估计距离步进（beam-safe —— 沿指向源的射线移动
    不会脱离定向源的覆盖波束）。若当前位置不可测（例如交会跳点落入定向盲区），
    则先退回已知可测的观测点。清除返回 True。"""
    b0, rel0 = sim.detect(sim.dog_pos[0], sim.dog_pos[1], ch)
    if (b0 is None or rel0 is None or rel0 <= 0.0) and obs:
        sim.move_to(*min(obs, key=lambda o: o[2])[0])
    pos = sim.dog_pos
    for _ in range(n_iter):
        if sim.try_clear(*pos):
            return True
        b, rel = sim.detect(pos[0], pos[1], ch)
        if b is None or rel is None or rel <= 0.0:
            return False
        step = _clamp(rel * R_REC_NOM, 15.0, max_step)
        pos = (pos[0] + step * math.cos(math.radians(b)),
               pos[1] + step * math.sin(math.radians(b)))
        sim.move_to(*pos)
        if sim.try_clear(*pos):
            return True
        obs.append((pos, b, rel))
    return sim.is_channel_cleared(ch)


def engage(sim, ch, obs, max_iter=6):
    """对频道 ch 的源执行「交会定位 → 归航逼近 → 光学/激光清除」，只用可观测量。

    顺序：① 已有 ≥2 观测 → 交会定位并跳至估计点；
          ② beam-safe 归航收敛（沿示向度逼近，不会脱离定向源波束）；
          ③ 归航失联 → 补一个正交第二观测点（交会角 90°）后再交会归航；
          ④ 仍失联 → 局部再捕获。返回 True 表示该频道源已清除。
    """
    for _ in range(max_iter):
        if sim.is_channel_cleared(ch):
            return True
        # ① 交会定位
        est = best_pair_fix(obs) if len(obs) >= 2 else None
        if est is not None:
            sim.move_to(*est)
            if sim.try_clear(*est):
                return True
            b3, rel3 = sim.detect(est[0], est[1], ch)
            if b3 is not None and rel3 is not None and rel3 > 0.0:
                obs.append((est, b3, rel3))
        # ② 归航
        if _homing(sim, ch, obs):
            return True
        # ③ 补正交第二观测点后再交会
        P1, b1, rel1 = min(obs, key=lambda o: o[2])
        d_est = _clamp(rel1 * R_REC_NOM, 80.0, R_REC_MAX * 0.9)
        P2 = perpendicular_second_point(P1, b1, d_est, offset=d_est)
        sim.move_to(*P2)
        b2, rel2 = sim.detect(P2[0], P2[1], ch)
        if b2 is not None and rel2 is not None and rel2 > 0.0:
            obs.append((P2, b2, rel2))
            est2 = cross_fix(P1, b1, P2, b2)
            if est2 is not None:
                sim.move_to(*est2)
                if sim.try_clear(*est2):
                    return True
                if _homing(sim, ch, obs):
                    return True
        # ④ 局部再捕获
        if local_hunt(sim, sim.dog_pos, ch, obs):
            return True
    return sim.is_channel_cleared(ch)


# ---------- 覆盖扫描路径 ----------
def _cover_ring_step(radius):
    """按环角色选取弧向步长 / m。

    * r ≤ 1000（内环）：COVER_STEP——环内任意点距环 ≤450 m，余量充足；
    * 1000 < r ≤ 1500（次外环）：COVER_STEP_MID——环内任意点距环 ≤450 m（仍远小于
      r_rec_min=1000 m），实测放大到 900 m 对全向与定向漏检率均无影响；
    * r > 1500（外侧环）：COVER_STEP_DIR——该环位于源分布环（R=1800 m）之外，
      朝外波束在此环上的角接受窗为 ±arccos(R_AREA/2000)=±25.8°，宽于环上相邻点
      半间距 15°，故可用最粗步长。数值验证见 coverage_detection_points。
    """
    if radius <= 1000.0:
        return COVER_STEP
    if radius <= 1500.0:
        return COVER_STEP_MID
    return COVER_STEP_DIR


def nn_route(points, start=(0.0, 0.0)):
    """贪心最近邻路线排序：从 start 出发每次取最近未访点。

    仅改变 list[(x,y)] 的顺序（签名与返回类型不变），用以压低覆盖扫描的
    点间行程；点规模 ≤ 80，O(n²) 足够。
    """
    rem = list(points)
    route = []
    cur = (float(start[0]), float(start[1]))
    while rem:
        j = min(range(len(rem)), key=lambda i: math.dist(cur, rem[i]))
        cur = rem[j]
        route.append(rem.pop(j))
    return route


def ring_detection_points(radius, step):
    """半径 radius 的整圆周检测点（相邻弧长 ≈ step）。"""
    n = max(8, int(round(2.0 * math.pi * radius / step)))
    return [(radius * math.cos(2.0 * math.pi * k / n),
             radius * math.sin(2.0 * math.pi * k / n)) for k in range(n)]


def coverage_detection_points(has_directional=False):
    """覆盖扫描检测点（同心环，已做最近邻路线排序）。

    圆域 R=1800 m、接收半径 r_rec≥1000 m：取环半径 {450, 1350}（=R·(2i−1)/(2·2)），
    则圆域内任意点到最近环的**径向**距离 ≤ 450 m。但完备性还需**弧向**项——环上相邻
    检测点的弧距有限，源可能落在两个检测点之间的弧段上。两项合并的确定性判据是
    「圆域内任意点到最近**检测点**的距离」，实测（ρ×θ 网格细扫）全向 ≤ 703.98 m
    （在 ρ=1800 m、θ=20° 处取得），仍远小于 r_rec_min=1000 m，故全向源覆盖完备。
    注意：只报径向 450 m 会漏掉弧向项，据此称「定理」属论证不完整——真正的保证
    来自「到最近检测点」的合成上界 703.98 m。

    定向源的外侧盲区：定向源只在其朝向 ±90° 内可测向。波束朝圆外的源**无法从
    源分布环（ρ≤1800 m）内侧测向**——可测要求检测点相对源的外向分量为正，即检测点
    半径 > 源半径，而圆域内不存在 ρ>1800 的点，故内侧环无论多密都只能逼近、无法覆盖
    贴边界的朝外源（这正是旧方案 {1200, 1799} 残余 2.5×10⁻⁴ 的几何根源）。
    因此定向覆盖另加一条位于源分布环**外侧**的环 r=2000 m：朝外波束在该环上的角
    接受窗为 ±arccos(R_AREA/2000)=±25.8°，宽于环上相邻点半间距 15°，故弧向步长可
    按角色分级放大到 COVER_STEP_DIR=1050 m（该环 12 点）。

    分级弧向步长：450 m 环 500 m（8 点）→ 1350 m 环 900 m（9 点）→ 2000 m 环 1050 m
    （12 点），合计 29 点（旧方案 {450,1350,1200,1799} 共 50 点）。

    ★ 数值验证：seed=7、20000 随机源下，全向 {450,1350} 漏检 0；定向 {450,1350,2000}
      漏检 0。另以 seed 1..40 × 20000、以及 4×10⁵ 独立样本复证，定向漏检均为 0。
    """
    radii = COVER_RADII_DIR if has_directional else COVER_RADII_OMNI
    pts = []
    for r in radii:
        pts += ring_detection_points(r, _cover_ring_step(r))
    return nn_route(pts)


def dog_strategy(sim, has_directional=False):
    """机器狗完整策略（测量驱动，无真值泄漏）：

    阶段A  覆盖扫描：沿同心环检测点逐点扫描未清除频道、记录示向度（含机会清除）；
    阶段B  逐频道定位清除：动态最近邻排序，对已发现频道执行「交会定位 → beam-safe
           归航逼近 → 光学/激光清除」，行程最短优先；
    阶段C  残余重试：对"已探测但未清除"的频道在环 r=900 m 上再扫描一次。

    has_directional=True（问题4）时阶段A额外包含一条位于源分布环**外侧**的环
    r=2000 m（COVER_RADII_DIR），用于从外侧覆盖朝外定向源（其内侧为测向盲区）。
    """
    obs = {ch: [] for ch in range(1, N_CHANNEL + 1)}
    det_time = {}
    trials = []

    def _record_clear(before):
        for sid in set(sim.cleared) - before:
            trials.append({"sid": sid, "time_s": sim.vtime})

    def _sweep(px, py):
        """在 (px,py) 扫描未清除频道并机会清除；返回发现信号的频道数。"""
        sim.move_to(px, py)
        while sim.try_clear(px, py):
            pass
        act = [ch for ch in range(1, N_CHANNEL + 1)
               if not sim.is_channel_cleared(ch)]
        if not act:
            return 0
        res = sim.sweep_all_channels(px, py, channels=act)
        for ch, (b, rel) in res.items():
            if ch not in det_time:
                det_time[ch] = sim.vtime
            obs[ch].append(((px, py), b, rel))
        return len(res)

    def _engage_pending():
        pending = [ch for ch in obs if obs[ch] and not sim.is_channel_cleared(ch)]

        def _anchor(ch):
            return min(obs[ch], key=lambda o: o[2])[0]

        while pending:
            ch = min(pending, key=lambda c: math.dist(sim.dog_pos, _anchor(c)))
            pending.remove(ch)
            if sim.is_channel_cleared(ch):
                continue
            before = set(sim.cleared)
            if engage(sim, ch, obs[ch]):
                _record_clear(before)

    # --- 阶段A：覆盖扫描 ---
    for (px, py) in coverage_detection_points(has_directional):
        _sweep(px, py)

    # --- 阶段B：逐频道定位清除（最近邻） ---
    _engage_pending()

    # --- 阶段C：残余重试（对已探测未清除的频道补扫环 r=900 m） ---
    if any(obs[ch] and not sim.is_channel_cleared(ch)
           for ch in range(1, N_CHANNEL + 1)):
        for (px, py) in nn_route(ring_detection_points(900.0, 400.0),
                                 sim.dog_pos):
            _sweep(px, py)
        _engage_pending()

    n_src = len(sim.sources)
    n_clear = len(sim.cleared)
    # 逐源定位清除耗时 = 清除时刻 − 该频道首次探测时刻
    per_src = []
    for src in sim.sources:
        ct = next((t["time_s"] for t in trials if t["sid"] == src.sid), None)
        dt = det_time.get(src.channel)
        if ct is not None and dt is not None:
            per_src.append(ct - dt)
    return {
        "n_sources": n_src,
        "n_cleared": n_clear,
        "cleared_fraction": n_clear / n_src if n_src else 1.0,
        "total_time_s": sim.vtime,
        "path_len_m": sim.path_len,
        "n_detect": sim.n_detect,
        "mean_time_per_source_s": sim.vtime / n_clear if n_clear else float("inf"),
        "mean_locate_clear_time_s": (float(np.mean(per_src)) if per_src
                                     else float("inf")),
        "max_locate_clear_time_s": (float(np.max(per_src)) if per_src
                                    else float("inf")),
        "uncleared_sids": sorted(s.sid for s in sim.sources
                                 if s.sid not in sim.cleared),
    }


def run_trials(n_trials=30, kind_mix=False, seed=SEED):
    """演练：n_trials 个随机案例，统计清除比例与定位清除时间。"""
    recs = []
    for t in range(n_trials):
        sim = Simulator(seed=seed + t, kind_mix=kind_mix)
        st = dog_strategy(sim, has_directional=kind_mix)
        st["seed"] = seed + t
        recs.append(st)
    fr = np.array([r["cleared_fraction"] for r in recs])
    tt = np.array([r["total_time_s"] for r in recs])
    lc = np.array([r["mean_locate_clear_time_s"] for r in recs])
    ps = np.array([r["mean_time_per_source_s"] for r in recs])
    return {
        "n_trials": n_trials,
        "mean_cleared_fraction": float(fr.mean()),
        "min_cleared_fraction": float(fr.min()),
        "all_cleared": bool((fr > 0.999).all()),
        "mean_total_time_s": float(tt.mean()),
        "std_total_time_s": float(tt.std()),
        "mean_locate_clear_time_s": float(lc.mean()),
        "std_locate_clear_time_s": float(lc.std()),
        "mean_time_per_source_s": float(ps.mean()),
    }, recs


# ----------------------------------------------------------------------
# 输出
# ----------------------------------------------------------------------
def _write_xlsx(path, r1, r2, recs3, recs4):
    from datetime import datetime, timezone
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "问题1_定位区域"
    ws.append(["案例", "交会角(°)", "顶点数", "直径(m)", "面积(m²)",
               "直径圆半径(m)", "最小包围圆半径(m)", "直径圆覆盖"])
    for k, v in r1.items():
        if "diameter_m" not in v:
            continue
        ws.append([k, v.get("crossing_angle_deg"), v.get("n_vertices"),
                   round(v["diameter_m"], 4), round(v.get("area_m2", 0.0), 4),
                   round(v["diam_circle_r_m"], 4), round(v["mec_r_m"], 4),
                   v["covers"]])
    ws2 = wb.create_sheet("问题2_候选点")
    ws2.append(["偏移(m)", "S2_x", "S2_y", "交会角(°)", "定位区域直径(m)",
                "定位区域面积(m²)"])
    for t in r2["tradeoff"]:
        ws2.append([t["offset_m"], round(t["S2"][0], 2), round(t["S2"][1], 2),
                    t["phi_deg"], round(t["region_diameter_m"], 4),
                    round(t["region_area_m2"], 4)])
    for sheet, recs in (("问题3_演练记录", recs3), ("问题4_演练记录", recs4)):
        w = wb.create_sheet(sheet)
        w.append(["seed", "源数", "清除数", "清除比例", "总时间(s)", "路径(m)",
                  "检测次数", "定位清除时间均值(s)"])
        for r in recs:
            w.append([r["seed"], r["n_sources"], r["n_cleared"],
                      round(r["cleared_fraction"], 4), round(r["total_time_s"], 2),
                      round(r["path_len_m"], 2), r["n_detect"],
                      round(r["mean_locate_clear_time_s"], 2)])
    # 确定性：openpyxl 默认把 wall-clock 写入 docProps/core.xml（created/modified），
    # 令同一输入产生不同字节。固定为注入时钟 MH_STATE_NOW，缺省用常量时间轴。
    _now = os.environ.get("MH_STATE_NOW")
    stamp = (datetime.fromisoformat(_now.replace("Z", "+00:00"))
             if _now else datetime(2026, 1, 1, tzinfo=timezone.utc))
    wb.properties.created = wb.properties.modified = stamp
    wb.save(path)
    # 二次确定性（openpyxl 3.1）：save_workbook 会把 properties.modified 重置为当前
    # 时钟（writer/excel.py:292），且经 zipfile 写盘时给每个 entry 打 wall-clock
    # date_time——两者都令字节随运行漂移。以黑盒方式重写 zip 容器：归一
    # core.xml 的 modified 字段 + 固定所有 entry 时间戳（同一输入 → 同一字节）。
    import re as _re
    import zipfile as _zip
    iso = stamp.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with _zip.ZipFile(path) as zf:
        members = [(i.filename, zf.read(i.filename)) for i in zf.infolist()]
    tmp = path + ".tmp"
    with _zip.ZipFile(tmp, "w", _zip.ZIP_DEFLATED) as zf:
        for name, blob in members:
            if name == "docProps/core.xml":
                blob = _re.sub(
                    rb"<dcterms:modified[^>]*>.*?</dcterms:modified>",
                    b'<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/"'
                    b' xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"'
                    b' xsi:type="dcterms:W3CDTF">' + iso + b"</dcterms:modified>",
                    blob)
            zi = _zip.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = _zip.ZIP_DEFLATED
            zf.writestr(zi, blob)
    os.replace(tmp, path)


def coverage_validation(n=20000, seed=7):
    """覆盖完备性蒙特卡洛验证：随机源（任意位置/半径/朝向）是否被检测点发现。

    全向源只要求距离 ≤ r_rec；定向源还要求落在定向方向 ±90° 内。
    """
    rng = random.Random(seed)

    def miss_ratio(points, omni):
        miss = 0
        for _ in range(n):
            rho = R_AREA * math.sqrt(rng.random())
            ang = rng.uniform(0.0, 2.0 * math.pi)
            x, y = rho * math.cos(ang), rho * math.sin(ang)
            rrec = rng.uniform(R_REC_MIN, R_REC_MAX)
            hd = None if omni else rng.uniform(0.0, 360.0)
            ok = False
            for (px, py) in points:
                if math.dist((px, py), (x, y)) <= rrec:
                    if omni:
                        ok = True
                        break
                    a = math.degrees(math.atan2(py - y, px - x)) % 360.0
                    if abs(((a - hd + 180.0) % 360.0) - 180.0) <= 90.0:
                        ok = True
                        break
            if not ok:
                miss += 1
        return miss / n

    return {
        "n_samples": n,
        "omni_miss_ratio": miss_ratio(coverage_detection_points(False), True),
        "dir_miss_ratio": miss_ratio(coverage_detection_points(True), False),
        "omni_cover_radii": [float(r) for r in COVER_RADII_OMNI],
        "dir_cover_radii": [float(r) for r in COVER_RADII_DIR],
    }


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    proj = os.path.abspath(os.path.join(here, "..", ".."))
    res_dir = os.path.join(proj, "artifacts", "results")
    os.makedirs(res_dir, exist_ok=True)
    random.seed(SEED)
    np.random.seed(SEED)

    r1 = problem1_demo()
    r2 = problem2_demo()

    print("=" * 68)
    print("问题3：全向干扰源 演练测试（30 个随机案例）")
    s3, recs3 = run_trials(n_trials=30, kind_mix=False)
    print(f"  清除比例 均值={s3['mean_cleared_fraction']:.4f} "
          f"最小={s3['min_cleared_fraction']:.4f} 全清除={s3['all_cleared']}")
    print(f"  平均定位清除时间={s3['mean_locate_clear_time_s']:.2f} s "
          f"(std {s3['std_locate_clear_time_s']:.2f})")
    print(f"  平均总时间={s3['mean_total_time_s']:.2f} s "
          f"(std {s3['std_total_time_s']:.2f})；"
          f"每源平均={s3['mean_time_per_source_s']:.2f} s")
    print("=" * 68)
    print("问题4：混合全向+定向 演练测试（30 个随机案例）")
    s4, recs4 = run_trials(n_trials=30, kind_mix=True)
    print(f"  清除比例 均值={s4['mean_cleared_fraction']:.4f} "
          f"最小={s4['min_cleared_fraction']:.4f} 全清除={s4['all_cleared']}")
    print(f"  平均定位清除时间={s4['mean_locate_clear_time_s']:.2f} s "
          f"(std {s4['std_locate_clear_time_s']:.2f})")
    print(f"  平均总时间={s4['mean_total_time_s']:.2f} s "
          f"(std {s4['std_total_time_s']:.2f})；"
          f"每源平均={s4['mean_time_per_source_s']:.2f} s")

    print("=" * 68)
    print("验证：覆盖完备性蒙特卡洛（各 20000 随机源）")
    val = coverage_validation()
    print(f"  全向源漏检率={val['omni_miss_ratio']:.6f}  "
          f"定向源漏检率={val['dir_miss_ratio']:.6f}")

    out = {
        "schema_version": 3,
        "problem": "CUMCM2026B",
        "random_seed": SEED,
        "model": "交会定位 + 覆盖搜索 + 逐源清除（本地模拟器演练）",
        "time_constants": {"detect": T_DETECT, "optic": T_OPTIC,
                           "laser": T_LASER, "speed_mps": V_DOG},
        "constants": {"R_AREA": R_AREA, "EPS_BEARING": EPS_BEARING,
                      "R_REC_MIN": R_REC_MIN, "R_REC_MAX": R_REC_MAX,
                      "N_CHANNEL": N_CHANNEL,
                      "cover_radii_omni": list(COVER_RADII_OMNI),
                      "cover_radii_dir": list(COVER_RADII_DIR),
                      "cover_step": COVER_STEP,
                      "cover_step_mid": COVER_STEP_MID,
                      "cover_step_dir": COVER_STEP_DIR},
        "summary": {"problem1": r1, "problem2": r2,
                    "problem3": s3, "problem4": s4},
        "validations": val,
    }
    # 结果聚合文件为多脚本共享：读-改-写（merge）保留其它脚本写入的段
    # （如 real_protocol_mock 由协议演练脚本并入），避免重跑本脚本抹掉他人产物。
    agg_path = os.path.join(proj, "all_results.json")
    merged = {}
    if os.path.exists(agg_path):
        try:
            with open(agg_path, encoding="utf-8") as f:
                merged = json.load(f)
        except (OSError, ValueError):
            merged = {}
    merged.update(out)
    with open(agg_path, "w", encoding="utf-8") as f:
        json.dump(merged, f, ensure_ascii=False, indent=2)
    _write_xlsx(os.path.join(res_dir, "resultB.xlsx"), r1, r2, recs3, recs4)

    print("\n[OK] all_results.json + resultB.xlsx 已写出")


if __name__ == "__main__":
    main()
